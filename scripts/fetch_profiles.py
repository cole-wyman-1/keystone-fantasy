"""Link each manager to a public profile and self-host a small headshot.

Sources, in priority order:
  1. keystone.com/our-people  — current employees. Each card has a stable /our-people/<slug> URL and a headshot.
  2. Cole's spreadsheets      — every Starter Files/private/*.xlsx except the master (git-ignored): Name, Link
                                (LinkedIn, or a keystone.com/our-people page for someone the matcher missed), and a
                                pasted headshot per row (optional; a keystone.com link gets the site's photo). Header row
                                optional. Images are read straight out of the .xlsx (xl/media) and tied to a row by
                                their drawing anchor.
  3. ALIASES below            — hand fixes for names that don't match automatically.

Writes data/profiles.json (slot id -> {url, source, photo}) and site/public/headshots/<slot>.jpg (160px, centre-cropped, path stored without a leading slash like logos.json).
compute.py attaches these to people.json. Run by hand after roster or spreadsheet changes; not part of the cron.
Idempotent: existing headshots are kept unless --force (so a vanished Keystone page or a re-saved spreadsheet
never blanks a photo by accident).
"""
from __future__ import annotations

import argparse
import difflib
import io
import json
import re
import sys
import unicodedata
import zipfile
from pathlib import Path

import requests
from PIL import Image, ImageOps

ROOT = Path(__file__).resolve().parent.parent
MASTER = ROOT / "data" / "master.json"
PRIVATE = ROOT / "Starter Files" / "private"
SHEETS = [p for p in sorted(PRIVATE.glob("*.xlsx")) if not p.name.startswith("Keystone_Fantasy_Football")]
OUT_JSON = ROOT / "data" / "profiles.json"
OUT_DIR = ROOT / "site" / "public" / "headshots"
PEOPLE_URL = "https://www.keystone.com/our-people"
SIZE = 160
UA = {"User-Agent": "Mozilla/5.0 (keystone-fantasy site builder)"}

# master display name -> keystone.com card name, for people the matcher can't pair on its own
ALIASES: dict[str, str] = {}

NICKNAMES = {
    "mike": "michael", "jeff": "jeffrey", "ben": "benjamin", "sam": "samuel", "chris": "christopher",
    "zach": "zachary", "max": "maxwell", "alex": "alexander", "matt": "matthew", "dan": "daniel",
    "danny": "daniel", "nick": "nicholas", "josh": "joshua", "jon": "jonathan", "andy": "andrew",
    "will": "william", "bill": "william", "tom": "thomas", "tim": "timothy", "jim": "james", "jake": "jacob",
    "rob": "robert", "bob": "robert", "ed": "edward", "ted": "edward", "katie": "katherine", "kate": "katherine",
    "cuau": "cuauhtemoc", "abby": "abigail", "liz": "elizabeth", "pat": "patrick", "greg": "gregory",
    "steve": "steven", "joe": "joseph", "tony": "anthony", "charlie": "charles",
}


def norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9 ]+", " ", s).strip()


def first_key(name: str) -> str:
    f = norm(name).split()[0] if norm(name) else ""
    return NICKNAMES.get(f, f)


def match(display: str, first: str, last: str, candidates: dict[str, str]) -> str | None:
    """candidates: normalized card name -> key. Exact full name, then nickname-aware first + exact last,
    then a close fuzzy match on the whole name (catches one-letter typos like Iakdawala/lakdawala)."""
    full = norm(display) or norm(f"{first} {last}")
    if full in candidates:
        return candidates[full]
    fk, lk = first_key(first or display), norm(last)
    hits = [k for n, k in candidates.items() if n.split()[-1] == lk and first_key(n) == fk]
    if len(hits) == 1:
        return hits[0]
    close = difflib.get_close_matches(full, list(candidates), n=1, cutoff=0.88)
    return candidates[close[0]] if close else None


# --- keystone.com -------------------------------------------------------------------------------
CARD_RE = re.compile(r'href="/our-people/([a-z0-9-]+)"[^>]*class="c-card-new[^"]*">(.*?)</a>', re.S)


def keystone_people(html: str) -> dict[str, dict]:
    """slug -> {name, photo}. The card's srcset carries a 500px rendition; fall back to the full image."""
    out = {}
    for slug, body in CARD_RE.findall(html):
        name = re.search(r'r-indexed="name"[^>]*>([^<]+)<', body)
        src = re.search(r'<img src="([^"]+)"', body)
        srcset = re.search(r'srcset="([^"]+)"', body)
        photo = None
        if srcset:
            for part in srcset.group(1).split(","):
                if part.strip().endswith(" 500w"):
                    photo = part.strip().split()[0]
        if not photo and src:
            photo = src.group(1)
        if name:
            out[slug] = {"name": name.group(1).strip(), "photo": photo}
    return out


# --- alumni spreadsheet ------------------------------------------------------------------------
def alumni_rows(path: Path) -> list[dict]:
    """[{name, url, image_bytes}] from the xlsx. Cells via openpyxl; pictures straight from the zip, each
    assigned to the sheet row its top edge sits in (an anchor whose rowOff is past half a row means
    'the next row' — Excel writes it that way when a picture is nudged to a row's top edge)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for i, r in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1):
        if i == 1 and not (len(r) > 1 and str(r[1] or "").startswith("http")):
            continue   # header row
        if r and r[0]:
            rows.append({"row": i, "name": str(r[0]).strip(), "url": (str(r[1]).strip() if len(r) > 1 and r[1] else ""), "image": None})
    by_row = {r["row"]: r for r in rows}

    z = zipfile.ZipFile(path)
    names = z.namelist()
    drawings = [n for n in names if re.match(r"xl/drawings/drawing\d+\.xml$", n)]
    for d in drawings:
        rels = z.read(d.replace("drawings/", "drawings/_rels/") + ".rels").decode()
        rid_to_file = {m.group(1): "xl/" + m.group(2).replace("../", "")
                       for m in re.finditer(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels)}
        xml = z.read(d).decode()
        for body in re.findall(r"<xdr:(?:one|two)CellAnchor[^>]*>(.*?)</xdr:(?:one|two)CellAnchor>", xml, re.S):
            fr = re.search(r"<xdr:from>.*?<xdr:row>(\d+)</xdr:row>\s*<xdr:rowOff>(-?\d+)</xdr:rowOff>", body, re.S)
            emb = re.search(r'r:embed="([^"]+)"', body)
            if not (fr and emb):
                continue
            row0, off = int(fr.group(1)), int(fr.group(2))
            sheet_row = row0 + 1 + (1 if off > 600_000 else 0)   # EMU; ~1.3M EMU per picture-height row here
            target = by_row.get(sheet_row)
            if target is None:
                print(f"  ! picture anchored at sheet row {sheet_row} has no name there; skipped", file=sys.stderr)
                continue
            if target["image"] is not None:
                print(f"  ! two pictures on row {sheet_row} ({target['name']}); keeping the first", file=sys.stderr)
                continue
            target["image"] = z.read(rid_to_file[emb.group(1)])
    return rows


# --- images -------------------------------------------------------------------------------------
def save_headshot(data: bytes, dest: Path) -> None:
    im = Image.open(io.BytesIO(data))
    im = ImageOps.exif_transpose(im).convert("RGB")
    im = ImageOps.fit(im, (SIZE, SIZE), Image.LANCZOS, centering=(0.5, 0.35))
    dest.parent.mkdir(parents=True, exist_ok=True)
    im.save(dest, "JPEG", quality=82, optimize=True)


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true", help="re-download / re-crop headshots that already exist")
    ap.add_argument("--offline", action="store_true", help="skip keystone.com; keep existing keystone entries")
    args = ap.parse_args()

    master = json.loads(MASTER.read_text())
    slots = [s for s in master["slots"] if not s["is_commissioner"]]
    previous = json.loads(OUT_JSON.read_text()) if OUT_JSON.exists() else {}
    profiles: dict[str, dict] = {}
    report = {"keystone": [], "linkedin": [], "unmatched_alumni_rows": [], "no_photo": [], "nobody": []}

    # 1. keystone.com
    cards: dict[str, dict] = {}
    if args.offline:
        profiles = {k: v for k, v in previous.items() if v.get("source") == "keystone"}
    else:
        html = requests.get(PEOPLE_URL, headers=UA, timeout=30).text
        cards = keystone_people(html)
        print(f"keystone.com: {len(cards)} people")
        cands = {norm(c["name"]): slug for slug, c in cards.items()}
        for s in slots:
            slug = cands.get(norm(ALIASES.get(s["display_name"], ""))) or match(s["display_name"], s["first_name"], s["last_name"], cands)
            if slug:
                profiles[s["slot_id"]] = {"url": f"{PEOPLE_URL}/{slug}", "source": "keystone", "name": cards[slug]["name"],
                                          "_photo_url": cards[slug]["photo"]}
                report["keystone"].append(f"{s['slot_id']} {s['display_name']} -> {cards[slug]['name']}")

    # 2. Cole's spreadsheets (only for people the keystone.com matcher didn't pair)
    cands = {norm(s["display_name"]): s["slot_id"] for s in slots}
    by_slot = {s["slot_id"]: s for s in slots}
    for sheet in SHEETS:
        rows = alumni_rows(sheet)
        print(f"{sheet.name}: {len(rows)} rows, {sum(1 for r in rows if r['image'])} photos")
        for r in rows:
            sid = match(r["name"], r["name"].split()[0], r["name"].split()[-1], cands)
            if not sid:
                report["unmatched_alumni_rows"].append(f"{r['name']} ({sheet.name})")
                continue
            if sid in profiles:
                print(f"  {r['name']} already matched ({profiles[sid]['source']}); row in {sheet.name} ignored")
                continue
            ks = re.match(r"https?://(?:www\.)?keystone\.com/our-people/([a-z0-9-]+)/?$", r["url"])
            if ks:
                card = cards.get(ks.group(1)) if not args.offline else None
                if not card and not args.offline:
                    print(f"  ! {r['name']}: keystone.com has no card for '{ks.group(1)}'; linking anyway", file=sys.stderr)
                profiles[sid] = {"url": f"{PEOPLE_URL}/{ks.group(1)}", "source": "keystone", "name": (card or {}).get("name", r["name"]),
                                 "_photo_url": (card or {}).get("photo"), "_image": r["image"]}
                report["keystone"].append(f"{sid} {by_slot[sid]['display_name']} -> {ks.group(1)} (from {sheet.name})")
            elif re.match(r"https?://(www\.)?linkedin\.com/", r["url"]):
                profiles[sid] = {"url": r["url"], "source": "linkedin", "name": r["name"], "_image": r["image"]}
                report["linkedin"].append(f"{sid} {by_slot[sid]['display_name']} -> {r['name']}")
            else:
                print(f"  ! {r['name']}: '{r['url']}' is neither LinkedIn nor keystone.com/our-people; skipped", file=sys.stderr)

    # 3. headshots
    for sid, p in profiles.items():
        dest = OUT_DIR / f"{sid.lower()}.jpg"
        if dest.exists() and not args.force:
            p["photo"] = f"headshots/{dest.name}"
        else:
            data = None
            if p.get("_image"):
                data = p["_image"]
            elif p.get("_photo_url") and not args.offline:
                try:
                    data = requests.get(p["_photo_url"], headers=UA, timeout=30).content
                except requests.RequestException as e:
                    print(f"  ! {sid}: headshot download failed ({e})", file=sys.stderr)
            if data:
                try:
                    save_headshot(data, dest)
                    p["photo"] = f"headshots/{dest.name}"
                except Exception as e:  # noqa: BLE001 - a bad image must not kill the run
                    print(f"  ! {sid}: could not decode headshot ({e})", file=sys.stderr)
        if not p.get("photo"):
            report["no_photo"].append(f"{sid} {p['name']}")
        p.pop("_image", None)
        p.pop("_photo_url", None)

    for s in slots:
        if s["slot_id"] not in profiles:
            report["nobody"].append(f"{s['slot_id']} {s['display_name']} ({s.get('keystone_group') or 'no group'})")

    OUT_JSON.write_text(json.dumps(dict(sorted(profiles.items())), indent=1, ensure_ascii=False) + "\n")
    print(f"\nwrote {OUT_JSON.relative_to(ROOT)}: {len(report['keystone'])} keystone, {len(report['linkedin'])} linkedin, "
          f"{len(report['nobody'])} with no profile")
    for key, title in (("unmatched_alumni_rows", "spreadsheet rows that match nobody"), ("no_photo", "profile but no photo"),
                       ("nobody", "no keystone.com page and not in any spreadsheet")):
        if report[key]:
            print(f"\n{title}:")
            for line in report[key]:
                print("  " + line)
    return 0


if __name__ == "__main__":
    sys.exit(main())
